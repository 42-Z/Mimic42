from __future__ import annotations

import asyncio
import logging
import random
from collections.abc import Awaitable, Callable, Mapping
from dataclasses import dataclass
from datetime import UTC, datetime
from enum import StrEnum
from typing import Any, Protocol, cast
from uuid import UUID, uuid4

from pydantic import BaseModel, Field
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from mimic42.core.activity import ActivityRecorder
from mimic42.core.album_grouper import AlbumGrouper
from mimic42.core.deferred_inbox import DeferredInbox
from mimic42.core.media import MAX_MEDIA_BYTES, MediaFile, MediaUploader
from mimic42.core.memory import MemoryServiceLike, RuntimeMemoryService
from mimic42.core.model_catalog import DEFAULT_LLM_MODEL
from mimic42.core.send_window import SendWindow, SendWindowTracker

logger = logging.getLogger("mimic42.agent_runtime")
logger.setLevel(logging.INFO)

# Дольше этого ждать открытия окна бессмысленно: сообщения успеют устареть.
DEFER_LIMIT_SECONDS = 300.0
# Слив ставится чуть позже открытия окна: иначе округление вниз запускает его
# раньше времени, и он перепланирует сам себя в плотном цикле.
DEFER_MARGIN_SECONDS = 0.25


class TelegramAuthorizationRequired(RuntimeError):
    """Raised when a Telethon user session is connected but not authorized."""


UNAUTHORIZED_SESSION_MESSAGE = (
    "Сессия Telegram не авторизована. Требуется повторная привязка Telegram-аккаунта."
)

REVOKED_SESSION_MESSAGE = (
    "Сессия Telegram недействительна. Требуется повторная привязка Telegram-аккаунта."
)


def _is_dead_session_error(exc: BaseException) -> bool:
    """Мёртвая сессия: нужен повторный вход, рантайм сам не восстановится.

    Telegram отдаёт это и как AuthKeyError (406, AUTH_KEY_DUPLICATED — обычная
    причина: одну сессию использовали с двух IP), и как UnauthorizedError
    (401: revoked/expired/unregistered/deactivated).
    """
    from telethon.errors import AuthKeyError, UnauthorizedError

    return isinstance(exc, (TelegramAuthorizationRequired, AuthKeyError, UnauthorizedError))


class AgentRuntimeState(StrEnum):
    STOPPED = "stopped"
    STARTING = "starting"
    RUNNING = "running"
    STOPPING = "stopping"
    ERROR = "error"


class AgentRuntimeConfig(BaseModel):
    agent_id: UUID
    owner_id: UUID
    telegram_api_id: int = Field(gt=0)
    telegram_api_hash: str = Field(min_length=1)
    telegram_session_string: str | None = Field(default=None, min_length=1)
    llm_model: str = Field(default=DEFAULT_LLM_MODEL, min_length=1)
    reasoning_effort: str = Field(default="high")
    system_prompt: str = Field(min_length=1)
    soul_prompt: str = Field(default="", max_length=20_000)
    name: str = Field(default="AI", min_length=1, max_length=120)

    @property
    def combined_prompt(self) -> str:
        prompt = self.system_prompt
        prompt = prompt.replace("{{name}}", self.name)
        prompt = prompt.replace("{{soul}}", self.soul_prompt)
        return prompt


class AgentStatus(BaseModel):
    agent_id: UUID
    owner_id: UUID
    state: AgentRuntimeState


class AgentTrigger(BaseModel):
    peer: str = Field(min_length=1)
    text: str = Field(min_length=1)
    raw_text: str = Field(default="")
    peer_name: str = Field(default="")
    chat_name: str = Field(default="")
    message_id: int | None = Field(default=None, gt=0)
    thread_id: UUID | None = None
    thread_title: str | None = None
    media: list[dict[str, Any]] = Field(default_factory=list)
    reply_to_message_id: int | None = Field(default=None, gt=0)
    reply_preview: str | None = None
    require_reply_to: bool = False
    """Ответ на схлопнутую пачку без reply_to нечитаем: рантайм подставит fallback."""
    fallback_reply_to: int | None = Field(default=None, gt=0)


class AgentTriggerResult(BaseModel):
    agent_id: UUID
    peer: str
    input_text: str
    response_text: str
    telegram_message_id: str | None = None


class TelegramClientLike(Protocol):
    async def __call__(self, request: Any) -> Any: ...

    async def connect(self) -> None: ...

    async def disconnect(self) -> None: ...

    async def is_user_authorized(self) -> bool: ...

    async def send_message(self, entity: str, message: str, **kwargs: Any) -> object: ...

    def add_event_handler(
        self,
        callback: Callable[[Any], Awaitable[None]],
        event: object | None = None,
    ) -> None: ...


class TelegramEventClientLike(Protocol):
    """The subset of the Telethon client surface used via incoming events."""

    async def __call__(self, request: Any) -> Any: ...

    async def download_media(self, message: Any, file: Any = None, **kwargs: Any) -> Any: ...


class TelegramEventLike(Protocol):
    """Structural type for Telethon NewMessage events used by the runtime."""

    chat_id: int | None
    sender_id: int | None
    client: TelegramEventClientLike

    async def get_chat(self) -> Any: ...

    async def get_reply_message(self) -> Any | None: ...

    async def get_input_chat(self) -> Any: ...


class LangChainAgentLike(Protocol):
    async def ainvoke(
        self,
        input_data: dict[str, object],
        context: object | None = None,
    ) -> object: ...


@dataclass(frozen=True)
class TurnContext:
    """Per-turn identity passed into the LangChain agent runtime context."""

    turn_id: str
    peer: str


@dataclass
class IncomingBlock:
    """Одно входящее (или один альбом), уже приведённое к тексту для модели."""

    text: str
    media: list[MediaFile]
    message_id: int | None
    reply_to_msg_id: int | None
    reply_preview: str
    raw_text: str
    thread_title: str | None
    sender_str: str
    chat_type_str: str
    msg_date: datetime


class MimicAgentRuntime:
    """Async runtime that owns one Telegram user session and one LangChain agent."""

    def __init__(
        self,
        *,
        config: AgentRuntimeConfig,
        telegram_client: TelegramClientLike,
        langchain_agent: LangChainAgentLike,
        memory_service: MemoryServiceLike | None = None,
        session_factory: async_sessionmaker[AsyncSession] | None = None,
        media_uploader: MediaUploader | None = None,
        send_window: SendWindowTracker | None = None,
    ) -> None:
        self.config = config
        self._telegram_client = telegram_client
        self._langchain_agent = langchain_agent
        self._memory_service = memory_service or RuntimeMemoryService()
        self._session_factory = session_factory
        self._media_uploader = media_uploader
        self._send_window = send_window
        self._state = AgentRuntimeState.STOPPED
        self._lifecycle_lock = asyncio.Lock()
        self._trigger_lock = asyncio.Lock()
        # Проверка окна и ход по ней — одна операция: иначе сообщения, пришедшие во время
        # хода, проходят гейт при ещё открытом окне и упираются в закрытое на отправке.
        self._dispatch_lock = asyncio.Lock()
        self._message_handler_registered = False
        self._member_tag_cache: dict[tuple[int, int], tuple[str | None, float]] = {}
        self._chat_mute_cache: dict[str, tuple[bool, float]] = {}
        self._scheduler_task: asyncio.Task[None] | None = None
        self._http_client: Any | None = None
        self._album_grouper = AlbumGrouper(self._flush_album)
        self._deferred_inbox = DeferredInbox(self._flush_deferred)
        self._activity = ActivityRecorder(session_factory) if session_factory is not None else None

    async def _record_event(
        self,
        *,
        event_type: str,
        status: str,
        payload: dict[str, Any] | None = None,
        error: str | None = None,
        started_at: datetime | None = None,
        completed_at: datetime | None = None,
    ) -> None:
        if self._activity is None:
            return
        await self._activity.record(
            agent_id=self.config.agent_id,
            event_type=event_type,
            status=status,
            payload=payload,
            error=error,
            started_at=started_at,
            completed_at=completed_at,
        )

    async def _mark_telegram_session_revoked(self, *, error: str) -> None:
        """Пометить telegram_sessions revoked: дэшборд предложит перепривязку.

        Ошибка записи не мешает основному исключению: статус агента и события
        фиксируются отдельно.
        """
        if self._session_factory is None:
            return
        try:
            from sqlalchemy import update
            from sqlalchemy.engine import CursorResult

            from mimic42.integrations.database_models import TelegramSessionModel

            async with self._session_factory() as db_session:
                result = await db_session.execute(
                    update(TelegramSessionModel)
                    .where(TelegramSessionModel.agent_id == self.config.agent_id)
                    .values(authorization_status="revoked", last_error=error)
                )
                await db_session.commit()
                # execute() статически возвращает Result, а rowcount есть только
                # у буферизованного CursorResult, который и приходит для UPDATE.
                if cast(CursorResult[Any], result).rowcount == 0:
                    logger.warning(
                        "No telegram_sessions row for agent %s, cannot mark revoked",
                        self.config.agent_id,
                    )
        except Exception:
            logger.warning(
                "Failed to mark telegram session revoked for agent %s",
                self.config.agent_id,
                exc_info=True,
            )

    @property
    def state(self) -> AgentRuntimeState:
        return self._state

    @property
    def status(self) -> AgentStatus:
        return AgentStatus(
            agent_id=self.config.agent_id,
            owner_id=self.config.owner_id,
            state=self.state,
        )

    async def start(self) -> None:
        async with self._lifecycle_lock:
            if self._state is AgentRuntimeState.RUNNING:
                logger.debug(f"Agent {self.config.agent_id} already running")
                return

            logger.info(f"Starting agent {self.config.agent_id}")
            self._state = AgentRuntimeState.STARTING
            try:
                logger.debug("Connecting to Telegram...")
                await self._telegram_client.connect()
                logger.debug("Connected. Checking authorization...")
                if not await self._telegram_client.is_user_authorized():
                    logger.error("Telegram session not authorized")
                    raise TelegramAuthorizationRequired(UNAUTHORIZED_SESSION_MESSAGE)
                logger.debug("Authorized. Registering message handler...")
                self._register_message_handler()
                logger.info("Message handler registered")
            except Exception as e:
                logger.error(f"Failed to start agent {self.config.agent_id}: {e}", exc_info=True)
                self._state = AgentRuntimeState.ERROR
                dead_session = _is_dead_session_error(e)
                reason = "unauthorized" if dead_session else "exception"
                if dead_session:
                    await self._mark_telegram_session_revoked(error=REVOKED_SESSION_MESSAGE)
                await self._record_event(
                    event_type="agent.start_failed",
                    status="failed",
                    payload={"reason": reason, "error_code": type(e).__name__},
                    error=REVOKED_SESSION_MESSAGE if dead_session else str(e),
                    started_at=datetime.now(UTC),
                    completed_at=datetime.now(UTC),
                )
                if dead_session and not isinstance(e, TelegramAuthorizationRequired):
                    # Унифицируем для вызывающего слоя: API отдаёт 428 с понятным
                    # русским текстом вместо 500. Исходное исключение остаётся
                    # в логе и в payload.error_code.
                    raise TelegramAuthorizationRequired(REVOKED_SESSION_MESSAGE) from e
                raise

            self._state = AgentRuntimeState.RUNNING
            if self._session_factory is not None:
                self._scheduler_task = asyncio.create_task(self._run_scheduler_loop())
            import httpx

            self._http_client = httpx.AsyncClient(timeout=30.0)
            await self._record_event(
                event_type="agent.started",
                status="succeeded",
                started_at=datetime.now(UTC),
                completed_at=datetime.now(UTC),
            )
            logger.info(f"Agent {self.config.agent_id} started successfully")

    async def stop(self) -> None:
        async with self._lifecycle_lock:
            if self._state is AgentRuntimeState.STOPPED:
                return

            self._state = AgentRuntimeState.STOPPING
            try:
                if self._scheduler_task is not None:
                    self._scheduler_task.cancel()
                    try:
                        await self._scheduler_task
                    except asyncio.CancelledError:
                        pass
                    self._scheduler_task = None

                await self._album_grouper.close()
                await self._deferred_inbox.close()

                if self._http_client is not None:
                    try:
                        await self._http_client.aclose()
                    finally:
                        self._http_client = None
            finally:
                # The Telegram client must be disconnected even if a previous
                # step failed — otherwise the userbot keeps answering messages
                # for a runtime nobody can reach anymore.
                await self._telegram_client.disconnect()
                self._state = AgentRuntimeState.STOPPED

            await self._record_event(
                event_type="agent.stopped",
                status="succeeded",
                started_at=datetime.now(UTC),
                completed_at=datetime.now(UTC),
            )

    async def close(self) -> None:
        """Stop for good and release what only a restart would need.

        The manager calls this when it drops the runtime (removal, reload,
        shutdown); a plain stop keeps the agent's HTTP client for a restart.
        """
        await self.stop()
        close_agent = getattr(self._langchain_agent, "aclose", None)
        if close_agent is not None:
            await close_agent()

    async def _humanized_send(
        self,
        peer: Any,
        text: str,
        reply_to: int | None = None,
    ) -> Any:
        """Send a message after a human-like typing delay with intermittent typing indicators."""
        from telethon import functions, types

        # Resolve peer entity for typing actions
        entity = peer
        get_input_entity = getattr(self._telegram_client, "get_input_entity", None)
        if callable(get_input_entity):
            try:
                entity = await get_input_entity(peer)
            except Exception:
                pass

        # Calculate typing delay: ~14 chars per second + random variance
        base_delay = len(text) * 0.07
        delay = max(0.8, min(15.0, base_delay + random.uniform(0.2, 1.2)))
        logger.debug("Humanized send: calculated delay %.2fs for %d chars", delay, len(text))

        end_time = asyncio.get_event_loop().time() + delay

        # Start typing
        try:
            await self._telegram_client(
                functions.messages.SetTypingRequest(
                    peer=entity,
                    action=types.SendMessageTypingAction(),
                )
            )
        except Exception:
            logger.debug("Failed to start typing action", exc_info=True)

        while True:
            remaining = end_time - asyncio.get_event_loop().time()
            if remaining <= 0:
                break
            sleep_for = min(0.5, remaining)
            await asyncio.sleep(sleep_for)

            # Small chance (8%) to briefly interrupt typing for realism
            if random.random() < 0.08:
                try:
                    await self._telegram_client(
                        functions.messages.SetTypingRequest(
                            peer=entity,
                            action=types.SendMessageCancelAction(),
                        )
                    )
                    await asyncio.sleep(random.uniform(0.05, 0.25))
                    await self._telegram_client(
                        functions.messages.SetTypingRequest(
                            peer=entity,
                            action=types.SendMessageTypingAction(),
                        )
                    )
                except Exception:
                    pass
            else:
                try:
                    await self._telegram_client(
                        functions.messages.SetTypingRequest(
                            peer=entity,
                            action=types.SendMessageTypingAction(),
                        )
                    )
                except Exception:
                    pass

        # Send the actual message
        sent_message = None
        try:
            # Telegram text limit is 4096 characters.
            # Use telethon.utils.split_text to respect markdown entities.
            if len(text) > 4096:
                from telethon.extensions import markdown as _md

                parsed_text, entities = _md.parse(text)
                from telethon import utils as _utils

                parts = list(_utils.split_text(parsed_text, entities, limit=4096))
                for i, (part_text, part_entities) in enumerate(parts):
                    part_reply_to = reply_to if i == 0 else None
                    sent_message = await self._telegram_client.send_message(
                        peer,
                        part_text,
                        formatting_entities=part_entities,
                        reply_to=part_reply_to,
                    )
            else:
                if reply_to is not None:
                    sent_message = await self._telegram_client.send_message(
                        peer,
                        text,
                        reply_to=reply_to,
                    )
                else:
                    sent_message = await self._telegram_client.send_message(peer, text)
        finally:
            # Cancel typing
            try:
                await self._telegram_client(
                    functions.messages.SetTypingRequest(
                        peer=entity,
                        action=types.SendMessageCancelAction(),
                    )
                )
            except Exception:
                pass

        return sent_message

    async def trigger_message(self, trigger: AgentTrigger) -> AgentTriggerResult:
        logger.debug(
            "trigger_message called: peer=%s, text=%s",
            trigger.peer,
            trigger.text[:50],
        )
        if self._state is not AgentRuntimeState.RUNNING:
            logger.info(f"Agent not running (state={self._state}), starting...")
            await self.start()

        async with self._trigger_lock:
            logger.debug(f"Processing message from {trigger.peer}: {trigger.text[:100]}")
            turn_id = str(uuid4())
            turn_context = TurnContext(turn_id=turn_id, peer=trigger.peer)
            reply_payload = (
                {
                    "message_id": trigger.reply_to_message_id,
                    "preview": trigger.reply_preview,
                }
                if trigger.reply_to_message_id is not None
                else None
            )
            messages = await self._memory_service.build_messages(
                agent_id=self.config.agent_id,
                peer=trigger.peer,
                user_text=trigger.text,
            )
            logger.debug(f"Built {len(messages)} messages for context")
            try:
                response = await self._langchain_agent.ainvoke(
                    {
                        "messages": messages,
                    },
                    context=turn_context,
                )
            except Exception as e:
                logger.error(f"Error invoking agent: {e}", exc_info=True)
                await self._record_event(
                    event_type="turn.failed",
                    status="failed",
                    payload={
                        "turn_id": turn_id,
                        "peer": trigger.peer,
                        "error_code": type(e).__name__,
                    },
                    error=str(e),
                    started_at=datetime.now(UTC),
                    completed_at=datetime.now(UTC),
                )
                # Keep the incoming message in the transcript even though the
                # turn crashed — the dashboard card must show what was asked.
                try:
                    await self._memory_service.save_messages(
                        agent_id=self.config.agent_id,
                        peer=trigger.peer,
                        input_messages=messages,
                        output_messages=[],
                        raw_user_text=trigger.raw_text,
                        turn_id=turn_id,
                        thread_id=trigger.thread_id,
                        media=trigger.media or None,
                        reply=reply_payload,
                    )
                except Exception:
                    logger.warning(
                        "Failed to persist incoming message after turn failure", exc_info=True
                    )
                # Mark the exception so the handler-level catch-all does not
                # record turn.failed a second time (the event above already
                # carries the turn_id).
                # Marker for the handler catch-all: turn.failed was already
                # recorded with the turn_id, so it must not be recorded twice.
                e._mimic_turn_failed_recorded = True  # ty: ignore[unresolved-attribute]
                raise

            output_messages = _messages_to_dicts(response)

            structured = _extract_structured_response(response)
            if structured is not None:
                send_any = bool(structured.get("send_any_message", True))
                response_text = structured.get("text", "")
                reply_to = structured.get("reply_to")
            else:
                send_any, response_text, reply_to = _interpret_agent_response(response)

            # Don't send empty messages
            if send_any and not response_text:
                logger.info("Agent generated empty response, not sending")
                send_any = False

            if send_any and trigger.require_reply_to and reply_to is None:
                reply_to = trigger.fallback_reply_to
                logger.info("Модель не указала reply_to на схлопнутой пачке, ставим %s", reply_to)

            # Между решением и отправкой прошло время генерации: слот мог закрыться.
            if send_any and self._send_window is not None:
                window = await self._send_window.check(trigger.peer)
                now = datetime.now(UTC)
                if not window.is_open(now):
                    logger.info(
                        "Окно отправки в %s закрыто (%s), ответ не уходит",
                        trigger.peer,
                        window.reason,
                    )
                    await self._record_event(
                        event_type="message.blocked",
                        status="cancelled",
                        payload={
                            "turn_id": turn_id,
                            "peer": trigger.peer,
                            "reason": window.reason,
                            "retry_after_seconds": window.retry_after(now),
                        },
                        started_at=now,
                        completed_at=datetime.now(UTC),
                    )
                    send_any = False

            # Convert stringified numeric peer ID to integer for Telethon compatibility
            peer_id_value: str | int = trigger.peer
            if isinstance(peer_id_value, str):
                if peer_id_value.startswith("-") and peer_id_value[1:].isdigit():
                    peer_id_value = int(peer_id_value)
                elif peer_id_value.isdigit():
                    peer_id_value = int(peer_id_value)
            # Use the correctly typed value for sending
            peer_id_for_send = peer_id_value
            # Mark incoming message as read immediately after deciding to reply,
            # before the typing delay, so the order is: read -> typing -> send.
            if send_any and trigger.message_id is not None:
                try:
                    from telethon.tl import functions

                    read_entity = peer_id_for_send
                    get_input_entity = getattr(self._telegram_client, "get_input_entity", None)
                    if callable(get_input_entity):
                        try:
                            read_entity = await get_input_entity(peer_id_for_send)
                        except Exception:
                            pass
                    await self._telegram_client(
                        functions.messages.ReadHistoryRequest(
                            peer=cast(Any, read_entity),
                            max_id=trigger.message_id,
                        )
                    )
                    logger.debug(
                        "Marked message %s as read in peer %s",
                        trigger.message_id,
                        trigger.peer,
                    )
                except Exception:
                    logger.warning("Failed to mark message as read", exc_info=True)

            sent_message = None
            if send_any:
                logger.info(f"Sending response to {peer_id_for_send}: {response_text[:100]}")
                try:
                    sent_message = await self._humanized_send(
                        peer_id_for_send,
                        response_text,
                        reply_to=reply_to,
                    )
                    logger.info(f"Message sent successfully to {peer_id_for_send}")
                    if self._send_window is not None:
                        self._send_window.note_sent(trigger.peer)
                except Exception as e:
                    if self._send_window is not None:
                        self._send_window.note_error(trigger.peer, e)
                    # The turn must not crash on a delivery failure, but the
                    # silence must be visible in the dashboard, not only in logs.
                    logger.exception("Failed to send Telegram message to %s", peer_id_for_send)
                    dead_session = _is_dead_session_error(e)
                    if dead_session:
                        await self._mark_telegram_session_revoked(error=REVOKED_SESSION_MESSAGE)
                    await self._record_event(
                        event_type="message.send_failed",
                        status="failed",
                        payload={
                            "turn_id": turn_id,
                            "peer": trigger.peer,
                            "error_code": type(e).__name__,
                        },
                        error=REVOKED_SESSION_MESSAGE if dead_session else str(e),
                        started_at=datetime.now(UTC),
                        completed_at=datetime.now(UTC),
                    )
            else:
                logger.debug("Agent decided not to send message (send_any=False)")

            await self._memory_service.save_messages(
                agent_id=self.config.agent_id,
                peer=trigger.peer,
                input_messages=messages,
                output_messages=output_messages,
                structured_response=structured,
                peer_name=trigger.peer_name,
                agent_name=self.config.name,
                raw_user_text=trigger.raw_text,
                turn_id=turn_id,
                thread_id=trigger.thread_id,
                media=trigger.media or None,
                reply=reply_payload,
            )

        return AgentTriggerResult(
            agent_id=self.config.agent_id,
            peer=trigger.peer,
            input_text=trigger.text,
            response_text=response_text,
            telegram_message_id=(
                _extract_message_id(sent_message) if sent_message is not None else None
            ),
        )

    async def _upsert_thread(
        self,
        *,
        peer: str,
        title: str | None,
        last_message_at: datetime | None,
    ) -> UUID | None:
        """Create or refresh the message_threads row for a peer.

        Failures are non-fatal: the thread exists for dashboard naming only.
        """
        if self._session_factory is None:
            return None
        try:
            from sqlalchemy import select

            from mimic42.integrations.database_models import MessageThreadModel

            async with self._session_factory() as db_session:
                thread = await db_session.scalar(
                    select(MessageThreadModel).where(
                        MessageThreadModel.agent_id == self.config.agent_id,
                        MessageThreadModel.telegram_peer_id == peer,
                    )
                )
                if thread is None:
                    thread = MessageThreadModel(
                        agent_id=self.config.agent_id,
                        telegram_peer_id=peer,
                        title=title,
                        last_message_at=last_message_at,
                    )
                    db_session.add(thread)
                    await db_session.flush()
                else:
                    if title:
                        thread.title = title
                    if last_message_at is not None:
                        thread.last_message_at = last_message_at
                await db_session.commit()
                return thread.id
        except Exception:
            logger.warning("Failed to upsert thread for peer %s", peer, exc_info=True)
            return None

    def _register_message_handler(self) -> None:
        if self._message_handler_registered:
            return

        try:
            from telethon import events
        except ImportError:
            event_builder = None
        else:
            event_builder = events.NewMessage(incoming=True)

        self._telegram_client.add_event_handler(self._handle_incoming_message, event_builder)
        self._message_handler_registered = True

    async def _handle_incoming_message(self, event: TelegramEventLike) -> None:
        """Альбомы буферизуются, одиночные сообщения обрабатываются сразу."""
        grouped_id = getattr(event, "grouped_id", None)
        if not isinstance(grouped_id, int):
            # TL: grouped_id — flags.17?long, то есть int или None. Любое другое
            # значение означает «это не элемент альбома».
            await self._dispatch_incoming([event])
            return
        chat_id = getattr(event, "chat_id", None)
        self._album_grouper.add((str(chat_id), str(grouped_id)), event)

    async def _flush_album(self, events: list[TelegramEventLike]) -> None:
        """Буфер альбома доставлен — обработать элементы одним ходом."""
        logger.info(
            "Grouped album with %d item(s) from chat %s",
            len(events),
            getattr(events[0], "chat_id", None),
        )
        try:
            await self._dispatch_incoming(events)
        except Exception:
            logger.exception("Failed to process grouped album")

    async def _is_chat_muted(self, event: TelegramEventLike, peer: str) -> bool:
        """Приглушён ли чат у самого агента (уведомления), а не запрет писать в него."""
        try:
            import time

            now_ts = time.time()

            is_muted = False
            if peer in self._chat_mute_cache:
                cached_muted, expiry = self._chat_mute_cache[peer]
                if now_ts < expiry:
                    is_muted = cached_muted
                    if is_muted:
                        return True
                else:
                    self._chat_mute_cache.pop(peer, None)

            if peer not in self._chat_mute_cache:
                input_chat = None
                get_input_chat = getattr(event, "get_input_chat", None)
                if callable(get_input_chat):
                    try:
                        input_chat = await event.get_input_chat()
                    except Exception:
                        logger.warning("Failed to get input chat for mute check", exc_info=True)

                if input_chat is None:
                    input_chat = getattr(event, "input_chat", None)

                if input_chat is None:
                    get_input_entity = getattr(self._telegram_client, "get_input_entity", None)
                    if callable(get_input_entity):
                        input_chat = await get_input_entity(peer)

                if input_chat is not None:
                    from telethon import functions, types

                    notify_peer = types.InputNotifyPeer(peer=input_chat)
                    res = await self._telegram_client(
                        functions.account.GetNotifySettingsRequest(peer=notify_peer)
                    )

                    is_muted = False
                    if getattr(res, "silent", False):
                        is_muted = True
                    if getattr(res, "mute_until", None):
                        from datetime import datetime as _datetime

                        if res.mute_until.tzinfo:
                            now = _datetime.now(res.mute_until.tzinfo)
                        else:
                            now = _datetime.now()
                        if res.mute_until > now:
                            is_muted = True

                    self._chat_mute_cache[peer] = (is_muted, now_ts + 60.0)

                    if is_muted:
                        logger.info("Chat %s is muted, skipping", peer)
                        return True
        except Exception:
            logger.exception("Failed to check mute status for peer %s", peer)

        return False

    async def _format_incoming(self, events: list[TelegramEventLike]) -> IncomingBlock | None:
        """Привести одно входящее или один альбом к тексту для модели."""
        event = events[0]
        # Элементы альбома обрабатываются по отдельности (у каждого свой
        # маркер и своя подпись), но ход, ответ и запись — общие.
        merged_content: list[str] = []
        merged_raw: list[str] = []
        media_files: list[MediaFile] = []
        for item in events:
            item_raw = getattr(item, "raw_text", None) or getattr(item, "text", None)
            if not isinstance(item_raw, str):
                item_raw = ""
            if item_raw:
                merged_raw.append(item_raw)
            item_text, item_media = await _process_media_and_text(
                item,
                item_raw,
                http_client=self._http_client,
                media_uploader=self._media_uploader,
                agent_id=self.config.agent_id,
            )
            if item_text:
                merged_content.append(item_text)
            media_files.extend(item_media)

        raw_text = "\n".join(merged_raw)
        text = "\n".join(merged_content)
        if not text:
            logger.info(
                "Empty text after _process_media_and_text for chat %s, skipping",
                getattr(event, "chat_id", None),
            )
            return None

        # Format sender name and metadata
        is_private = getattr(event, "is_private", False)
        is_group = getattr(event, "is_group", False)

        from datetime import datetime

        msg_date = getattr(event, "date", None)
        if not msg_date:
            message = getattr(event, "message", None)
            msg_date = getattr(message, "date", None)
        if not msg_date:
            msg_date = datetime.now()
        time_str = msg_date.strftime("%Y-%m-%d %H:%M:%S")

        # Chat name
        chat = None
        if is_private:
            chat_type_str = "ЛС"
        else:
            chat = await event.get_chat()
            chat_title = getattr(chat, "title", "")
            if not chat_title:
                chat_title = getattr(chat, "username", "") or str(getattr(event, "chat_id", ""))
            if is_group:
                chat_type_str = f'Группа "{chat_title}"'
            else:
                chat_type_str = f'Канал "{chat_title}"'

        # Sender details
        get_sender = getattr(event, "get_sender", None)
        sender = None
        if callable(get_sender):
            try:
                import inspect

                res = get_sender()
                if inspect.isawaitable(res):
                    sender = await res
                else:
                    sender = res
            except Exception:
                logger.warning("Failed to get sender for event", exc_info=True)
        if sender:
            first_name = getattr(sender, "first_name", None) or ""
            last_name = getattr(sender, "last_name", None) or ""
            name_parts = []
            if first_name:
                name_parts.append(first_name)
            if last_name:
                name_parts.append(last_name)
            name_str = " ".join(name_parts)
            if not name_str:
                name_str = (
                    getattr(sender, "title", None)
                    or getattr(sender, "username", None)
                    or str(getattr(sender, "id", ""))
                )
            if not name_str:
                name_str = "Unknown"

            username = getattr(sender, "username", None)
            username_str = f"@{username}" if username else ""
            sender_id = getattr(sender, "id", None)
            id_str = f"ID: {sender_id}" if sender_id else ""

            details = ", ".join(filter(None, [username_str, id_str]))
            details_str = f" ({details})" if details else ""
            sender_str = f"{name_str}{details_str}"
        else:
            chat = await event.get_chat()
            if isinstance(chat, str):
                sender_str = chat
            else:
                chat_title = getattr(chat, "title", None)
                sender_str = chat_title if isinstance(chat_title, str) else "Unknown"

        # Check role/title
        title = None
        event_chat_id = event.chat_id
        event_sender_id = event.sender_id
        if event_sender_id and event_chat_id:
            cache_key = (event_chat_id, event_sender_id)
            import time

            now_ts = time.time()
            if cache_key in self._member_tag_cache:
                cached_title, expiry = self._member_tag_cache[cache_key]
                if now_ts < expiry:
                    title = cached_title

            is_expired = (
                cache_key not in self._member_tag_cache
                or now_ts >= self._member_tag_cache[cache_key][1]
            )
            if is_expired:
                try:
                    from telethon.tl import functions

                    is_supergroup = getattr(event, "is_channel", False)
                    if is_supergroup:
                        input_chat = getattr(event, "input_chat", None) or event_chat_id
                        input_sender = getattr(event, "input_sender", None) or event_sender_id
                        res = await event.client(
                            functions.channels.GetParticipantRequest(
                                channel=cast(Any, input_chat),
                                participant=cast(Any, input_sender),
                            )
                        )
                        title = res.participant.title if hasattr(res.participant, "title") else None
                except Exception:
                    logger.warning("Failed to get participant title", exc_info=True)
                    title = None
                self._member_tag_cache[cache_key] = (title, now_ts + 3600.0)

        # Fallback to channel post author signature
        post_author = getattr(getattr(event, "message", None), "post_author", None)
        if not title and post_author:
            title = post_author

        if title:
            sender_str += f" [Подпись/Роль: {title}]"

        # Thread title for the dashboard: reuse the entities already
        # fetched above — no extra Telegram requests. For private chats
        # the interlocutor is the chat itself.
        thread_title: str | None = None
        thread_entity = chat if (not is_private and chat is not None) else sender
        if thread_entity is not None:
            from telethon import utils as telethon_utils

            thread_title = telethon_utils.get_display_name(thread_entity) or None
        if not thread_title:
            thread_title = str(getattr(event, "chat_id", "")) or None

        # NewMessage.Event delegates __getattr__ to self.message, but
        # self.message is a raw types.Message (not the custom wrapper),
        # so it lacks the reply_to_msg_id property. Inspect reply_to directly.
        from telethon.tl import types

        reply_to_msg_id = None
        ev_message = getattr(event, "message", None)
        if ev_message:
            reply_to = getattr(ev_message, "reply_to", None)
            if isinstance(reply_to, types.MessageReplyHeader):
                reply_to_msg_id = reply_to.reply_to_msg_id
            elif reply_to:
                reply_to_msg_id = getattr(reply_to, "reply_to_msg_id", None)

        reply_str = ""
        reply_preview = ""
        if reply_to_msg_id:
            logger.debug(
                "Reply detected: reply_to_msg_id=%s for chat_id=%s",
                reply_to_msg_id,
                getattr(event, "chat_id", None),
            )
            reply_preview = ""
            try:
                reply_msg = await event.get_reply_message()
                if reply_msg:
                    raw = getattr(reply_msg, "raw_text", "")
                    reply_preview = raw or getattr(reply_msg, "text", "")
                    logger.debug(
                        "get_reply_message succeeded, preview=%s",
                        reply_preview[:30] if reply_preview else "(empty)",
                    )
                else:
                    logger.debug("get_reply_message returned None")
            except Exception:
                logger.debug("get_reply_message failed", exc_info=True)

            if not reply_preview:
                # Fallback: fetch the replied message directly via RPC
                try:
                    get_messages = getattr(self._telegram_client, "get_messages", None)
                    if callable(get_messages):
                        peer = await _extract_incoming_peer(event)
                        msgs = await get_messages(peer, ids=reply_to_msg_id)
                        if msgs:
                            reply_msg = msgs[0] if isinstance(msgs, list) else msgs
                            raw = getattr(reply_msg, "raw_text", "")
                            reply_preview = raw or getattr(reply_msg, "text", "")
                            logger.debug(
                                "Fallback get_messages succeeded, preview=%s",
                                reply_preview[:30] if reply_preview else "(empty)",
                            )
                except Exception:
                    logger.debug("Fallback get_messages failed", exc_info=True)

            if reply_preview:
                preview = reply_preview[:20]
                reply_str = f'Ответ на сообщение #{reply_to_msg_id} ("{preview}...")\n'
            else:
                reply_str = f"Ответ на сообщение #{reply_to_msg_id}\n"

        # Format output message text
        incoming_msg_id = _extract_incoming_message_id(event)
        album_note = ""
        if len(events) > 1:
            item_ids = [
                str(item_id)
                for item in events
                if (item_id := _extract_incoming_message_id(item)) is not None
            ]
            album_note = f"Альбом из {len(events)} файлов (ID: {', '.join(item_ids)})\n"
        text = (
            f"[Входящее сообщение]\n"
            f"Время: {time_str}\n"
            f"Чат: {chat_type_str}\n"
            f"Отправитель: {sender_str}\n"
            f"ID сообщения: {incoming_msg_id}\n"
            f"{album_note}"
            f"{reply_str}"
            f"Содержимое: {text}"
        )

        return IncomingBlock(
            text=text,
            media=media_files,
            message_id=_extract_incoming_message_id(event),
            reply_to_msg_id=reply_to_msg_id,
            reply_preview=reply_preview,
            raw_text=raw_text,
            thread_title=thread_title,
            sender_str=sender_str,
            chat_type_str=chat_type_str,
            msg_date=msg_date,
        )

    async def _dispatch_incoming(self, events: list[TelegramEventLike]) -> None:
        """Решить, идёт ли ход сейчас, позже или не идёт вовсе."""
        event = events[0]
        logger.info("Incoming message event received")
        logger.info(
            "Incoming message event: chat_id=%s, text=%s",
            getattr(event, "chat_id", None),
            getattr(event, "raw_text", "")[:50],
        )
        try:
            peer = await _extract_incoming_peer(event)
        except Exception as e:
            await self._record_incoming_failure(str(getattr(event, "chat_id", "") or ""), e)
            return
        if await self._is_chat_muted(event, peer):
            return
        async with self._dispatch_lock:
            await self._gate_and_process(event, events, peer)

    async def _gate_and_process(
        self, event: TelegramEventLike, events: list[TelegramEventLike], peer: str
    ) -> None:
        # Окно отправки бывает только у групп. В ЛС ограничений на запись нет, а пост
        # канала читают, не отвечая в него: комментарий уходит в связанную группу.
        if self._send_window is None or not getattr(event, "is_group", False):
            await self._process_batch(peer, [events])
            return

        chat = None
        get_chat = getattr(event, "get_chat", None)
        if callable(get_chat):
            try:
                # event.chat бывает пустым: Telegram не всегда шлёт эти данные.
                chat = await get_chat()
            except Exception:
                logger.warning("Не удалось получить чат для проверки окна", exc_info=True)

        now = datetime.now(UTC)
        window = await self._send_window.check(peer, chat=chat)
        if window.is_open(now):
            # Сбрасываем объявление: следующее закрытие снова надо сообщить.
            self._send_window.announce(peer, "open")
            await self._process_batch(peer, [events])
            return

        retry_after = window.retry_after(now)
        if retry_after is not None and retry_after <= DEFER_LIMIT_SECONDS:
            self._deferred_inbox.add(peer, list(events), delay=_delay_until(window, now))
            if self._send_window.announce(peer, window.reason):
                await self._record_event(
                    event_type="message.deferred",
                    status="succeeded",
                    payload={
                        "peer": peer,
                        "reason": window.reason,
                        "retry_after_seconds": retry_after,
                    },
                    started_at=now,
                    completed_at=datetime.now(UTC),
                )
            return

        # Закрыто надолго или бессрочно: копить нечего, сообщения устареют раньше.
        if self._send_window.announce(peer, window.reason):
            await self._record_event(
                event_type="message.write_forbidden",
                status="cancelled",
                payload={
                    "peer": peer,
                    "reason": window.reason,
                    "until": window.open_at.isoformat() if window.open_at else None,
                },
                started_at=now,
                completed_at=datetime.now(UTC),
            )
            await self._notify_write_forbidden(peer, window)

    async def _notify_write_forbidden(self, peer: str, window: SendWindow) -> None:
        """Один служебный ход: агент узнаёт про запрет и может отреагировать иначе."""
        until = (
            f" до {window.open_at:%Y-%m-%d %H:%M}"
            if window.open_at is not None and not window.forever
            else ""
        )
        try:
            await self.trigger_message(
                AgentTrigger(
                    peer=peer,
                    text=(
                        "[Системное уведомление]\n"
                        f"В чате {peer} у тебя забрали право писать{until}. "
                        "Отправить туда ничего не получится — ни ответом, ни инструментом. "
                        "Входящие оттуда ты больше не увидишь, пока запрет не снимут."
                    ),
                )
            )
        except Exception as e:
            await self._record_incoming_failure(peer, e)

    async def _flush_deferred(
        self, peer: str, groups: list[list[Any]], arrived: list[float]
    ) -> None:
        """Окно должно было открыться: перепроверяем и разбираем накопленное одним ходом."""
        async with self._dispatch_lock:
            # Слив в полёте stop() не отменяет, а trigger_message на остановленном
            # рантайме запустил бы его заново и ответил бы в чат после остановки.
            if self._state is not AgentRuntimeState.RUNNING:
                logger.info("Рантайм остановлен, отложенное из чата %s отброшено", peer)
                return
            if self._send_window is not None:
                now = datetime.now(UTC)
                window = await self._send_window.check(peer)
                if not window.is_open(now):
                    retry_after = window.retry_after(now)
                    if retry_after is not None and retry_after <= DEFER_LIMIT_SECONDS:
                        # Слот успел закрыться снова (свой ответ инструментом, ошибка Telegram).
                        delay = max(_delay_until(window, now), 1.0)
                        for added_at, group in zip(arrived, groups, strict=True):
                            self._deferred_inbox.add(peer, group, delay=delay, added_at=added_at)
                    else:
                        logger.info("Окно в чате %s закрыто надолго, накопленное отброшено", peer)
                    return
            await self._process_batch(peer, groups)

    async def _process_batch(self, peer: str, groups: list[list[TelegramEventLike]]) -> None:
        """Один ход по нескольким группам входящих (сообщение или альбом)."""
        # Защищаем остальной конвейер обработки сообщения от падений.
        try:
            blocks: list[IncomingBlock] = []
            for group in groups:
                block = await self._format_incoming(group)
                if block is not None:
                    blocks.append(block)
            if not blocks:
                logger.info("Пачка для чата %s пуста после разбора, пропускаем", peer)
                return

            last = blocks[-1]
            body = "\n\n".join(block.text for block in blocks)
            window = await self._send_window.check(peer) if self._send_window else None
            text = f"{_batch_header(len(blocks), window)}{body}"
            media: list[dict[str, Any]] = []
            for block in blocks:
                media.extend(m.as_payload() for m in block.media)

            thread_id = await self._upsert_thread(
                peer=peer,
                title=last.thread_title,
                last_message_at=last.msg_date,
            )
            await self.trigger_message(
                AgentTrigger(
                    peer=peer,
                    text=text,
                    raw_text="\n".join(block.raw_text for block in blocks),
                    peer_name=last.sender_str,
                    chat_name=last.chat_type_str,
                    message_id=last.message_id,
                    thread_id=thread_id,
                    thread_title=last.thread_title,
                    media=media,
                    reply_to_message_id=last.reply_to_msg_id,
                    reply_preview=last.reply_preview[:200] if last.reply_preview else None,
                    require_reply_to=len(blocks) > 1,
                    fallback_reply_to=last.message_id,
                )
            )
        except Exception as e:
            await self._record_incoming_failure(peer, e)

    async def _record_incoming_failure(self, peer: str, exc: Exception) -> None:
        logger.error("Unhandled exception in incoming message handler", exc_info=exc)
        if getattr(exc, "_mimic_turn_failed_recorded", False):
            # trigger_message already recorded turn.failed with the
            # turn_id — a second row would double the error KPI.
            return
        await self._record_event(
            event_type="turn.failed",
            status="failed",
            payload={"peer": peer, "error_code": type(exc).__name__},
            error=str(exc),
            started_at=datetime.now(UTC),
            completed_at=datetime.now(UTC),
        )

    async def _run_scheduler_loop(self) -> None:
        """Background loop to check and trigger pending agent timers."""
        while self._state is AgentRuntimeState.RUNNING:
            try:
                await self._check_and_trigger_timers()
            except asyncio.CancelledError:
                break
            except Exception:
                logger.exception("Error in scheduler loop")

            try:
                await asyncio.sleep(5.0)
            except asyncio.CancelledError:
                break

    async def _check_and_trigger_timers(self) -> None:
        if not self._session_factory:
            return

        from datetime import UTC, datetime

        from sqlalchemy import select, update

        from mimic42.integrations.database_models import AgentTimerModel

        now_utc = datetime.now(UTC)

        async with self._session_factory() as db_session:
            stmt = select(AgentTimerModel).where(
                AgentTimerModel.agent_id == self.config.agent_id,
                AgentTimerModel.status == "pending",
                AgentTimerModel.trigger_at <= now_utc,
            )
            result = await db_session.scalars(stmt)
            due_timers = list(result)

            if not due_timers:
                return

            for timer in due_timers:
                timer.status = "running"
            await db_session.commit()

            for timer in due_timers:
                try:
                    trigger_text = f"[Отложенное событие] Сработал таймер: {timer.description}"
                    await self.trigger_message(
                        AgentTrigger(
                            peer=timer.peer,
                            text=trigger_text,
                        )
                    )
                    timer.status = "succeeded"
                    await self._record_event(
                        event_type="timer.fired",
                        status="succeeded",
                        payload={
                            "peer": timer.peer,
                            "description": timer.description,
                            "trigger_at": timer.trigger_at.isoformat(),
                        },
                    )
                except Exception as e:
                    logger.exception("Error triggering timer %s", timer.id)
                    timer.status = "failed"
                    await self._record_event(
                        event_type="timer.failed",
                        status="failed",
                        payload={
                            "peer": timer.peer,
                            "description": timer.description,
                            "trigger_at": timer.trigger_at.isoformat(),
                            "error_code": type(e).__name__,
                        },
                        error=str(e),
                    )

                # Update status
                async with self._session_factory() as update_session:
                    await update_session.execute(
                        update(AgentTimerModel)
                        .where(AgentTimerModel.id == timer.id)
                        .values(status=timer.status)
                    )
                    await update_session.commit()


def _delay_until(window: SendWindow, now: datetime) -> float:
    """Секунды до открытия окна с небольшим запасом (см. DEFER_MARGIN_SECONDS)."""
    if window.open_at is None:
        return DEFER_MARGIN_SECONDS
    return max(0.0, (window.open_at - now).total_seconds()) + DEFER_MARGIN_SECONDS


def _batch_header(block_count: int, window: SendWindow | None) -> str:
    """Шапка над входящим: сколько накопилось и сколько стоит ответ.

    Формулировка подобрана замером на самой слабой модели каталога. Прежняя
    («ответить можно только ОДНИМ сообщением») читалась как требование
    отвечать: в чужом разговоре модель отвечала в 4 из 4 случаев против 2 из 4
    без шапки. Теперь ответ явно необязателен и привязан к адресованности,
    а пропущенный reply_to достраивает рантайм (fallback_reply_to).
    """
    seconds = window.slowmode_seconds if window is not None else None
    if block_count <= 1:
        if seconds is None:
            return ""
        return (
            f"[В чате медленный режим: одно сообщение раз в {seconds} с. "
            f"Ответишь — следующее сможешь написать не раньше чем через {seconds} с]\n\n"
        )
    mode = f" Медленный режим: одно сообщение раз в {seconds} с." if seconds is not None else ""
    return (
        f"[Накопилось сообщений: {block_count}.{mode} "
        "Ответить можно один раз и только если что-то из этого адресовано тебе — тогда "
        "укажи в reply_to ID нужного сообщения. Иначе send_any_message = false]\n\n"
    )


def _sticker_file_of(message: Any) -> tuple[str, str]:
    """Имя и mime файла стикера: статичный — webp, анимированные — tgs/webm.
    Иначе анимированный стикер получает битый mime и не открывается в ленте."""
    doc = getattr(getattr(message, "media", None), "document", None)
    mime = getattr(doc, "mime_type", None)
    if mime == "application/x-tgsticker":
        return "sticker.tgs", "application/x-tgsticker"
    if isinstance(mime, str) and "webm" in mime:
        return "sticker.webm", "video/webm"
    return "sticker.webp", "image/webp"


async def _process_media_and_text(
    event: TelegramEventLike,
    text: str,
    *,
    http_client: Any | None = None,
    media_uploader: MediaUploader | None = None,
    agent_id: UUID | None = None,
) -> tuple[str, list[MediaFile]]:
    message = getattr(event, "message", None)
    if not message or not getattr(message, "media", None):
        return text, []

    media_files: list[MediaFile] = []

    async def _archive(kind: str, filename: str, mime_type: str, data: bytes) -> None:
        """Archive one attachment to Storage; never break the turn on failure."""
        if media_uploader is None or agent_id is None or not data:
            return
        try:
            archived = await media_uploader.upload(
                agent_id=agent_id,
                filename=filename,
                data=data,
                mime_type=mime_type,
                kind=kind,
            )
        except Exception:
            logger.warning("Media archiving failed for %s", filename, exc_info=True)
            return
        if archived is not None:
            media_files.append(archived)

    try:
        from mimic42.integrations.telegram_tools import format_media_object

        media_id = format_media_object(message)
        if not media_id:
            return text, media_files

        if media_id.startswith("photo:"):
            data = await event.client.download_media(message, file=bytes)
            await _archive("photo", "photo.jpeg", "image/jpeg", data or b"")
            return f"[Фото id={media_id}]" + (f" {text}" if text else ""), media_files

        elif media_id.startswith("sticker:"):
            parts = media_id.split(":")
            emoji = parts[5] if len(parts) > 5 else ""
            pack_name = parts[6] if len(parts) > 6 else ""
            pack_str = f" пак={pack_name}" if pack_name else ""
            data = await event.client.download_media(message, file=bytes)
            sticker_name, sticker_mime = _sticker_file_of(message)
            await _archive("sticker", sticker_name, sticker_mime, data or b"")
            return (
                f"[Стикер {emoji} id={media_id}{pack_str}]" + (f" {text}" if text else ""),
                media_files,
            )

        elif media_id.startswith(("voice:", "round:")):
            from io import BytesIO

            import httpx

            from mimic42.config import Settings

            settings = Settings()
            api_key = settings.openrouter_api_key
            if not api_key:
                err_msg = "[Голосовое сообщение (ошибка: OPENROUTER_API_KEY не установлен)]"
                return err_msg + (f" {text}" if text else ""), media_files

            buffer = BytesIO()
            await event.client.download_media(message, file=buffer)
            file_bytes = buffer.getvalue()
            if not file_bytes:
                err_msg = "[Голосовое сообщение (ошибка: файл пустой)]"
                return err_msg + (f" {text}" if text else ""), media_files

            filename = "voice.ogg" if media_id.startswith("voice:") else "video.mp4"
            is_voice = media_id.startswith("voice:")
            await _archive(
                "voice" if is_voice else "round",
                filename,
                "audio/ogg" if is_voice else "video/mp4",
                file_bytes,
            )

            try:
                client = http_client
                if client is None:
                    client = httpx.AsyncClient(timeout=30.0)
                import base64

                audio_b64 = base64.b64encode(file_bytes).decode("utf-8")
                payload = {
                    "model": "openai/whisper-large-v3",
                    "input_audio": {
                        "data": audio_b64,
                        "format": filename.split(".")[-1],
                    },
                }
                headers = {
                    "Authorization": f"Bearer {api_key}",
                    "Content-Type": "application/json",
                }
                response = await client.post(
                    "https://openrouter.ai/api/v1/audio/transcriptions",
                    headers=headers,
                    json=payload,
                )
                response.raise_for_status()
                res_json = response.json()
                transcription = res_json.get("text", "")
                mtype = "Голосовое сообщение" if media_id.startswith("voice:") else "Видеосообщение"
                trans_text = f'[{mtype} (расшифровка: "{transcription}")]'
                return trans_text + (f" {text}" if text else ""), media_files
            except Exception as e:
                mtype = "Голосовое сообщение" if media_id.startswith("voice:") else "Видеосообщение"
                return (
                    f"[{mtype} (ошибка транскрипции: {e})]" + (f" {text}" if text else ""),
                    media_files,
                )

        elif media_id.startswith("doc:"):
            parts = media_id.split(":")
            filename = parts[5] if len(parts) > 5 else "file"
            ext = filename.split(".")[-1].lower() if "." in filename else ""

            allowed_exts = (
                "docx",
                "xlsx",
                "txt",
                "md",
                "json",
                "csv",
                "xml",
                "py",
                "html",
                "css",
                "yaml",
                "yml",
            )
            from io import BytesIO

            media_obj = getattr(message, "media", None)
            doc_obj = getattr(media_obj, "document", None)
            doc_mime = getattr(doc_obj, "mime_type", None)
            doc_mime_type = doc_mime if isinstance(doc_mime, str) else "application/octet-stream"
            doc_size = getattr(doc_obj, "size", None)

            if ext not in allowed_exts and ext != "":
                # The LLM cannot read this type, but the dashboard must still be
                # able to open the file from the logs — archive it (with the
                # same size cap as the storage layer) unless it is huge.
                if not isinstance(doc_size, int) or doc_size <= MAX_MEDIA_BYTES:
                    buffer = BytesIO()
                    await event.client.download_media(message, file=buffer)
                    await _archive("doc", filename, doc_mime_type, buffer.getvalue())
                return (
                    f"[Файл name={filename} (этот тип документа нельзя открыть)]"
                    + (f" {text}" if text else ""),
                    media_files,
                )

            # Size cap applies to readable types too: the file is pulled into
            # memory and fed to the LLM, so a huge one must not be downloaded.
            if isinstance(doc_size, int) and doc_size > MAX_MEDIA_BYTES:
                return (
                    f"[Файл name={filename} (слишком большой: {doc_size} байт)]"
                    + (f" {text}" if text else ""),
                    media_files,
                )

            buffer = BytesIO()
            await event.client.download_media(message, file=buffer)
            file_bytes = buffer.getvalue()
            if not file_bytes:
                return (
                    f"[Файл name={filename} (пустой)]" + (f" {text}" if text else ""),
                    media_files,
                )

            await _archive("doc", filename, doc_mime_type, file_bytes)

            if ext == "docx":
                try:
                    import docx

                    doc = docx.Document(BytesIO(file_bytes))
                    paragraphs = [p.text for p in doc.paragraphs]
                    for table in doc.tables:
                        for row in table.rows:
                            row_text = [cell.text for cell in row.cells]
                            paragraphs.append(" | ".join(row_text))
                    doc_content = "\n".join(paragraphs)
                    doc_text = f'[Файл name={filename} (содержимое: "{doc_content}")]'
                    return doc_text + (f" {text}" if text else ""), media_files
                except Exception as e:
                    err_msg = f"[Файл name={filename} (ошибка чтения: {e})]"
                    return err_msg + (f" {text}" if text else ""), media_files

            elif ext == "xlsx":
                try:
                    import openpyxl

                    wb = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True)
                    sheet_texts = []
                    for sheet in wb.worksheets:
                        sheet_texts.append(f"Лист: {sheet.title}")
                        for row in sheet.iter_rows(values_only=True):
                            row_str = " | ".join(str(val) if val is not None else "" for val in row)
                            if row_str.strip(" |"):
                                sheet_texts.append(row_str)
                    xlsx_content = "\n".join(sheet_texts)
                    xlsx_text = f'[Файл name={filename} (содержимое: "{xlsx_content}")]'
                    return xlsx_text + (f" {text}" if text else ""), media_files
                except Exception as e:
                    err_msg = f"[Файл name={filename} (ошибка чтения: {e})]"
                    return err_msg + (f" {text}" if text else ""), media_files

            else:
                try:
                    txt_content = file_bytes.decode("utf-8")
                except UnicodeDecodeError:
                    try:
                        txt_content = file_bytes.decode("latin-1")
                    except Exception as e:
                        txt_content = f"<Ошибка декодирования: {e}>"

                txt_text = f'[Файл name={filename} (содержимое: "{txt_content}")]'
                return txt_text + (f" {text}" if text else ""), media_files

    except Exception:
        pass

    return text, media_files


def _messages_to_dicts(response: object) -> list[dict[str, Any]]:
    """Extract and normalize the messages list from a LangGraph agent response."""
    if isinstance(response, Mapping):
        resp_map = cast("Mapping[str, Any]", response)
        raw_messages = resp_map.get("messages")
        if isinstance(raw_messages, list):
            return [_message_to_dict(m) for m in raw_messages]
    return []


def _message_to_dict(msg: object) -> dict[str, Any]:
    """Convert a LangChain BaseMessage (or dict) to a normalized dict."""
    if isinstance(msg, Mapping):
        return dict(cast("Mapping[str, Any]", msg))
    if hasattr(msg, "model_dump"):
        model_dump = getattr(msg, "model_dump", None)
        if callable(model_dump):
            return model_dump()
    # Fallback for plain objects with attributes
    result: dict[str, Any] = {}
    for attr in ("type", "role", "content", "tool_calls", "tool_call_id", "id", "name"):
        val = getattr(msg, attr, None)
        if val is not None:
            result[attr] = val
    return result


def _extract_response_text(response: object) -> str:
    if isinstance(response, str):
        return response
    if isinstance(response, Mapping):
        response_map = cast("Mapping[str, Any]", response)
        messages = response_map.get("messages")
        if isinstance(messages, list) and messages:
            # Find the last AIMessage (assistant response), not just the last message
            # (which could be a ToolMessage if tools were called).
            # BaseMessage objects use `type` ("ai", "human", "tool"), not `role`.
            for message in reversed(messages):
                if isinstance(message, Mapping):
                    msg_map = cast("Mapping[str, Any]", message)
                    role = msg_map.get("role")
                    if role == "assistant":
                        content = _get_content(message)
                        if content:
                            return content
                else:
                    # Check for BaseMessage objects (AIMessage has type="ai")
                    msg_type = getattr(message, "type", None)
                    role = getattr(message, "role", None)
                    if role == "assistant" or msg_type == "ai":
                        content = _get_content(message)
                        if content:
                            return content
        output = response_map.get("output") or response_map.get("content")
        if isinstance(output, str) and output:
            return output
    content = _get_content(response)
    if content:
        return content
    # Fallback: return empty string instead of str(response) to avoid sending garbage
    logger.warning("Could not extract response text from: %s", type(response))
    return ""


def _get_content(value: object) -> str | None:
    if isinstance(value, Mapping):
        value_map = cast("Mapping[str, Any]", value)
        content = value_map.get("content")
    else:
        content = getattr(value, "content", None)
    if isinstance(content, str) and content:
        return content
    if isinstance(content, list):
        parts: list[str] = []
        for item in content:
            if isinstance(item, str):
                parts.append(item)
            elif isinstance(item, Mapping):
                item_map = cast("Mapping[str, Any]", item)
                text = item_map.get("text")
                if isinstance(text, str):
                    parts.append(text)
        if parts:
            return "\n".join(parts)
    return None


def _extract_structured_response(response: object) -> dict[str, Any] | None:
    """Extract the validated structured response from a LangChain agent result.

    When ``create_agent`` is configured with ``response_format``, the final
    state dict contains a ``structured_response`` key holding the parsed
    schema instance (Pydantic model, dataclass, dict, etc.).

    Returns the response as a plain dict, or ``None`` if absent / unsupported.
    """
    if not isinstance(response, Mapping):
        return None
    resp_map = cast("Mapping[str, Any]", response)
    structured = resp_map.get("structured_response")
    if structured is None:
        return None
    if isinstance(structured, BaseModel):
        return structured.model_dump()
    if isinstance(structured, Mapping):
        return dict(structured)
    return None


def _interpret_agent_response(response: object) -> tuple[bool, str, int | None]:
    """Interpret a LangChain agent response into (send_any_message, text, reply_to).

    This function supports:
    - plain string responses
    - mapping-like responses produced by LangChain (dict-like), including parsed
      output from OutputParsers / PydanticOutputParser where fields like
      'send_any_message', 'text', and 'reply_to' may be present.

    Falls back to extracting the best available textual content.
    """
    # Default behaviour: send message, text from _extract_response_text, no reply
    default_text = _extract_response_text(response)

    if isinstance(response, Mapping):
        resp_map = cast("Mapping[str, Any]", response)
        # send_any_message may be absent; treat non-bool as True
        raw_send = resp_map.get("send_any_message")
        send_any = bool(raw_send) if isinstance(raw_send, bool) else True

        # prefer explicit 'text' field, then 'output'/'content' via extractor
        txt = resp_map.get("text")
        if isinstance(txt, str) and txt:
            text = txt
        else:
            text = default_text

        reply_val = resp_map.get("reply_to") or resp_map.get("reply_to_message_id")
        reply_to = None
        if isinstance(reply_val, int):
            reply_to = reply_val
        elif isinstance(reply_val, str) and reply_val.isdigit():
            try:
                reply_to = int(reply_val)
            except Exception:
                reply_to = None

        return send_any, text, reply_to

    # Non-mapping responses: send and use extracted text
    return True, default_text, None


def _extract_message_id(message: object) -> str | None:
    if isinstance(message, Mapping):
        message_map = cast("Mapping[str, Any]", message)
        value: Any = message_map.get("id")
    else:
        value = getattr(message, "id", None)
    if value is None:
        return None
    return str(value)


def _extract_incoming_text(event: object) -> str:
    text = getattr(event, "raw_text", None) or getattr(event, "text", None)
    if not isinstance(text, str):
        text = ""

    message = getattr(event, "message", None)
    if message and getattr(message, "media", None):
        try:
            from mimic42.integrations.telegram_tools import format_media_object

            media_id = format_media_object(message)
            if media_id:
                if media_id.startswith("photo:"):
                    text = f"[Фото id={media_id}]" + (f" {text}" if text else "")
                elif media_id.startswith("sticker:"):
                    parts = media_id.split(":")
                    emoji = parts[5] if len(parts) > 5 else ""
                    pack_name = parts[6] if len(parts) > 6 else ""
                    pack_str = f" пак={pack_name}" if pack_name else ""
                    text = f"[Стикер {emoji} id={media_id}{pack_str}]" + (
                        f" {text}" if text else ""
                    )
                elif media_id.startswith("voice:"):
                    text = f"[Голосовое сообщение id={media_id}]" + (f" {text}" if text else "")
                elif media_id.startswith("round:"):
                    text = f"[Видеосообщение id={media_id}]" + (f" {text}" if text else "")
                elif media_id.startswith("doc:"):
                    parts = media_id.split(":")
                    filename = parts[5] if len(parts) > 5 else "file"
                    text = f"[Файл name={filename} id={media_id}]" + (f" {text}" if text else "")
        except Exception:
            pass

    return text


async def _extract_incoming_peer(event: object) -> str:
    chat_id = getattr(event, "chat_id", None)
    if chat_id is not None:
        return str(chat_id)
    get_chat = getattr(event, "get_chat", None)
    if callable(get_chat):
        chat = await get_chat()
        if chat:
            cid = getattr(chat, "id", None)
            if cid is not None:
                return str(cid)
    peer_id = getattr(event, "peer_id", None)
    if peer_id is not None:
        return str(peer_id)
    raise ValueError("Incoming Telegram event does not include a peer")


def _extract_incoming_message_id(event: object) -> int | None:
    value = getattr(event, "id", None)
    if isinstance(value, int):
        return value
    message = getattr(event, "message", None)
    message_id = getattr(message, "id", None)
    if isinstance(message_id, int):
        return message_id
    return None
